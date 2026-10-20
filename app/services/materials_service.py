"""课程材料：把本地文件夹里的课件、讲义、试卷读成**带来源标记**的文本。

为什么单独一层，以及为什么"带来源标记"是硬要求：
    资料模式与"纯 AI 编排"的全部差别在这里——后者的知识来自模型的先验，
    前者的知识必须来自用户给的这些文件。所以解析结果要能回答
    "这段文字出自哪个文件的哪一页"：
      · 讲师引用资料时才能核对；
      · 我（以及用户）才可能发现"讲的内容资料里根本没有"；
      · 章节才能映射回"来自哪份材料的哪几页"。

只做文本提取，不走模型视觉（一期决定）：
    PDF/PPT 通常是电子生成的，里面的文字可以直接取出来。
    把它们转成图片再让模型"看"，成本会高一个数量级，而且更慢。
    扫描件（整页是图片、取不到文字）**会在这里被如实标记出来**，
    而不是悄悄当成"这份材料没内容"——那会让人以为资料本身是空的。
"""

from __future__ import annotations

import math
import re
import threading
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from ..materials import MODE_STRICT, normalize_mode

# 支持的材料类型。刻意只收"文字能直接取出来"的格式。
TEXT_SUFFIXES = {".txt", ".md"}
PDF_SUFFIXES = {".pdf"}
PPTX_SUFFIXES = {".pptx"}
DOCX_SUFFIXES = {".docx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | PDF_SUFFIXES | PPTX_SUFFIXES | DOCX_SUFFIXES

# 单个文件最多读进多少字符。资料是用来"编排"和"引用"的，不是全量背诵，
# 超长的部分截断即可——但**截断这件事要留痕**（见 truncated 标记）。
MAX_CHARS_PER_FILE = 60000
# 扫描件判定：一页提取出来的文字少于这个数，基本可以认为没取到文字
SCAN_HINT_CHARS = 12

# 一次取原文最多给讲师几块。
#
# 为什么要有个上限：`sources` 有时指得比较粗（例如"整份讲义"），
# 而一份讲义可能有几十页。全给出去就是**上下文污染**——与项目其它地方
# （只给局部大纲、只给最近几轮对话）是同一个理由：这一轮用不着的，
# 不只是浪费字数，还会把讲师该讲的重点稀释掉。
MAX_EXCERPT_BLOCKS = 6

_WS = re.compile(r"[ \t\u3000]+")
_BLANKS = re.compile(r"\n{3,}")


@dataclass
class Block:
    """一段有出处的文本。"""

    source: str          # 文件名（不含目录），用于回指
    locator: str         # "第 3 页" / "第 5 张幻灯片" / "全文"
    text: str


@dataclass
class FileReport:
    """单个文件的解析结果与"发生了什么"。"""

    path: Path
    blocks: list[Block] = field(default_factory=list)
    error: str = ""            # 解析失败的原因（空 = 成功）
    scanned_pages: int = 0     # 取不到文字的页数（多半是扫描件/图片页）
    truncated: bool = False    # 是否因为超长被截断

    @property
    def ok(self) -> bool:
        return not self.error

    @property
    def has_text(self) -> bool:
        return any(b.text.strip() for b in self.blocks)


def _clean(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = _WS.sub(" ", text)
    return _BLANKS.sub("\n\n", text).strip()


def list_materials(folder: Path, *, recursive: bool = True) -> list[Path]:
    """列出文件夹里受支持的材料文件，按路径排序。

    排序只是为了**有个确定的初始顺序**（否则同一批文件每次读到的次序都可能不同，
    结果不可复现）。它**不代表**资料的正确顺序——资料的顺序问题
    由编排阶段按内容判断，用户也可以在问卷里说明。
    """
    if not folder.is_dir():
        return []
    pattern = "**/*" if recursive else "*"
    found = [
        p for p in folder.glob(pattern)
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$")
    ]
    return sorted(found, key=lambda p: str(p).lower())


def _read_pdf(path: Path) -> FileReport:
    import fitz  # PyMuPDF

    report = FileReport(path=path)
    with fitz.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            text = _clean(page.get_text("text"))
            if len(text) < SCAN_HINT_CHARS:
                report.scanned_pages += 1
                continue
            report.blocks.append(Block(path.name, f"第 {index} 页", text))
    return report


def _read_pptx(path: Path) -> FileReport:
    from pptx import Presentation

    report = FileReport(path=path)
    presentation = Presentation(str(path))
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            # 文本框与表格都能取到文字；图片取不到——那属于"没文字"，
            # 由下面的 scanned_pages 记一笔，不做特殊处理。
            if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                for paragraph in shape.text_frame.paragraphs:
                    line = _clean("".join(run.text for run in paragraph.runs))
                    if line:
                        parts.append(line)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [_clean(cell.text) for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))

        text = _clean("\n".join(parts))
        if not text:
            report.scanned_pages += 1
            continue
        report.blocks.append(Block(path.name, f"第 {index} 张幻灯片", text))
    return report


def _read_docx(path: Path) -> FileReport:
    import docx

    report = FileReport(path=path)
    document = docx.Document(str(path))
    paragraphs = [_clean(p.text) for p in document.paragraphs]
    text = _clean("\n".join(p for p in paragraphs if p))

    if not text:
        # 可能是整篇都是图片的 docx
        report.scanned_pages = 1
        return report
    report.blocks.append(Block(path.name, "全文", text))
    return report


def _read_text(path: Path) -> FileReport:
    report = FileReport(path=path)
    # errors="replace"：资料里混进个别坏字节时，宁可看到几个问号，
    # 也不要整份文件读不进来。
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = _clean(raw)
    if text:
        report.blocks.append(Block(path.name, "全文", text))
    return report


_READERS = {
    **{s: _read_pdf for s in PDF_SUFFIXES},
    **{s: _read_pptx for s in PPTX_SUFFIXES},
    **{s: _read_docx for s in DOCX_SUFFIXES},
    **{s: _read_text for s in TEXT_SUFFIXES},
}


def parse_file(path: Path) -> FileReport:
    """解析单个文件。**绝不抛异常**——一份坏文件不该让整批资料读不进来。"""
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        return FileReport(path=path, error=f"暂不支持的类型：{path.suffix}")

    try:
        report = reader(path)
    except Exception as exc:  # noqa: BLE001 —— 原因要如实报给用户
        return FileReport(path=path, error=f"{exc.__class__.__name__}: {exc}")

    total = sum(len(b.text) for b in report.blocks)
    if total > MAX_CHARS_PER_FILE:
        report.truncated = True
        kept: list[Block] = []
        used = 0
        for block in report.blocks:
            if used >= MAX_CHARS_PER_FILE:
                break
            room = MAX_CHARS_PER_FILE - used
            if len(block.text) > room:
                kept.append(Block(block.source, block.locator, block.text[:room]))
                used = MAX_CHARS_PER_FILE
            else:
                kept.append(block)
                used += len(block.text)
        report.blocks = kept
    return report


def parse_folder(folder: Path, *, limit: int = 0) -> list[FileReport]:
    """解析整个文件夹。`limit>0` 时只取前 N 个（用于先试跑）。"""
    paths = list_materials(folder)
    if limit > 0:
        paths = paths[:limit]
    return [parse_file(p) for p in paths]


# 解析结果按"目录 + 文件指纹"缓存。
#
# 为什么必须有：生成大纲那一次读盘是几分钟的一次性成本，而**讲师的每一轮**都要取原文。
# 几百 MB 的课件每轮重读一遍，会把"讲一句就出来"变成"先等十秒"——
# 实测一门课的课件常有几十份 PDF，这一项足以毁掉流畅感。
# 指纹用文件集合 + 大小 + mtime：用户往文件夹里补了新课件，缓存自然失效，
# 不必让他手动点一次"重新解析"。
_REPORTS_CACHE: OrderedDict[tuple, list[FileReport]] = OrderedDict()
_REPORTS_CACHE_LOCK = threading.Lock()
# 缓存几个目录：单用户本地用，够在几门课之间来回切就行。
_REPORTS_CACHE_MAX = 4


def folder_fingerprint(folder: Path) -> tuple:
    """目录内容的指纹：文件名 + 大小 + 修改时间。

    只 stat 不读取——这正是它便宜的原因（相对重新解析一遍 PDF）。
    """
    entries = []
    for path in list_materials(folder):
        try:
            stat = path.stat()
        except OSError:
            # 权限不足/文件正被占用：不进指纹，解析时按失败如实上报
            continue
        entries.append((str(path).lower(), stat.st_size, int(stat.st_mtime)))
    return (str(folder), tuple(entries))


def load_reports(folder: Path) -> list[FileReport]:
    """解析文件夹（带缓存）——讲师每一轮取原文都走这里。

    **返回值请只读**：它是缓存里的同一个对象，就地修改会污染其它轮次。
    """
    key = folder_fingerprint(folder)

    with _REPORTS_CACHE_LOCK:
        cached = _REPORTS_CACHE.get(key)
        if cached is not None:
            _REPORTS_CACHE.move_to_end(key)
            return cached

    reports = parse_folder(folder)

    with _REPORTS_CACHE_LOCK:
        _REPORTS_CACHE[key] = reports
        _REPORTS_CACHE.move_to_end(key)
        while len(_REPORTS_CACHE) > _REPORTS_CACHE_MAX:
            _REPORTS_CACHE.popitem(last=False)
    return reports


def clear_report_cache() -> None:
    """丢掉解析缓存。

    "重新生成大纲"时用它：那是一次**显式的重新开始**，
    用户的意图是"资料变了，按新的来"，不该让缓存挡住。
    """
    with _REPORTS_CACHE_LOCK:
        _REPORTS_CACHE.clear()


def render_digest(reports: list[FileReport], *, per_file_limit: int = 1500) -> str:
    """把解析结果压成一段可以放进提示词的文本。

    每个文件给一段**开头**（足以看出它在讲什么、怎么组织），
    而不是全文——全量塞进上下文既贵又容易让模型抓不住重点。
    真正的原文在讲师那一轮按知识点取用。
    """
    lines: list[str] = []
    for report in reports:
        name = report.path.name
        if not report.ok:
            lines.append(f"### {name}\n（读取失败：{report.error}）")
            continue
        if not report.has_text:
            lines.append(f"### {name}\n（没有取到文字，可能是扫描件或纯图片）")
            continue

        head = "\n".join(b.text for b in report.blocks)[:per_file_limit]
        note = []
        if report.truncated:
            note.append("内容较长，此处只给开头")
        if report.scanned_pages:
            note.append(f"有 {report.scanned_pages} 处没取到文字（可能是图片）")
        suffix = f"（{'；'.join(note)}）" if note else ""
        lines.append(f"### {name}{suffix}\n{head}")
    return "\n\n".join(lines)


def summarize_reports(reports: list[FileReport]) -> dict:
    """给界面用的一句统计，也用于日志。"""
    ok = [r for r in reports if r.ok and r.has_text]
    failed = [r for r in reports if not r.ok]
    empty = [r for r in reports if r.ok and not r.has_text]
    return {
        "total": len(reports),
        "parsed": len(ok),
        "failed": len(failed),
        "no_text": len(empty),
        "failed_names": [r.path.name for r in failed],
        "no_text_names": [r.path.name for r in empty],
    }


def summarize_materials(
    reports: list[FileReport], *, on_progress: Callable[[int, int, str], None] | None = None
) -> tuple[list[dict], list[str]]:
    """逐份材料读成结构化摘要，返回 (摘要列表, 失败说明列表)。

    **单份失败不影响其余**：一份格式古怪的材料不该让整门课的编排卡住。
    失败的记下文件名与原因，交回给界面如实说明——用户有权知道
    "哪几份没读进去"，否则他会以为资料已经被完整参考过了。

    `on_progress(序号, 总数, 文件名)` 用来把"正在读第几份"推给界面：
    这一步可能要好几分钟，没有进度的话用户只能干等，会以为卡死了。
    """
    from ..agents.material_reader import read_material  # 函数内导入：避开启动期依赖

    usable = [r for r in reports if r.ok and r.has_text]
    digests: list[dict] = []
    failures: list[str] = []
    total = len(usable)

    for index, report in enumerate(usable, start=1):
        if on_progress is not None:
            on_progress(index, total, report.path.name)
        text = "\n".join(block.text for block in report.blocks)
        try:
            digests.append(read_material(filename=report.path.name, content=text, total=total))
        except Exception as exc:  # noqa: BLE001 —— 原因要如实带回给用户
            failures.append(f"{report.path.name}：{exc}")

    return digests, failures


def render_digests(digests: list[dict], failures: list[str], *, structure_limit: int = 12,
                   points_limit: int = 8) -> str:
    """把各份材料的摘要拼成一段"整批资料的概览"，供编排使用。

    这里给的是**摘要**而不是原文：编排要看清的是资料的整体轮廓
    （有哪些材料、各自讲什么、按什么顺序），逐字原文留给讲师那一轮按需取用。
    """
    lines: list[str] = []
    for digest in digests:
        lines.append(f"### {digest.get('file', '（未命名）')}")
        if digest.get("role"):
            lines.append(f"- 角色：{digest['role']}")
        if digest.get("topic"):
            lines.append(f"- 主题：{digest['topic']}")
        structure = (digest.get("structure") or [])[:structure_limit]
        if structure:
            lines.append("- 它的组织结构（照原样）：")
            lines.extend(f"  {i}. {item}" for i, item in enumerate(structure, start=1))
        points = (digest.get("key_points") or [])[:points_limit]
        if points:
            lines.append("- 重点：")
            lines.extend(f"  - {item}" for item in points)
        lines.append("")

    if failures:
        lines.append("### 没能读进去的材料")
        lines.extend(f"- {item}" for item in failures)
        lines.append("")

    return "\n".join(lines).strip()


def _blocks_for(report: FileReport, locator: str, *, strict: bool = False) -> list[Block]:
    """在一个文件里按位置找块。

    **严格模式下不做"退回整个文件"**（评审 4.1）：
        原来的三级退化最后一级是"都对不上就给我整份文件"。这在参考模式下
        是合理兜底，但在严格模式下会变成**静默换料**——学生指定"第 2 页"，
        系统拿来的是整份讲义，而讲师仍被要求"严格依据资料"，
        于是它可能引用到学生没指定的页。那正好破坏了严格模式唯一的价值：可核验。

    所以 `strict=True` 时只有前两级（精确匹配、页号范围匹配）；
    定位不上就返回空，由上层如实告诉用户"没取到"，而不是换一份材料。
    唯一的例外是 `locator` 为空——那表示"这个来源只指到文件级"，取全文是符合意图的。
    """
    if not locator:
        return list(report.blocks)

    exact = [b for b in report.blocks if b.locator == locator]
    if exact:
        return exact

    # 位置常见的写法是范围（"第 2-3 页"）：把数字抠出来按页号匹配。
    # 只在 locator 里找数字，不去块正文里找——正文里阿拉伯数字太多，会误命中。
    wanted = {int(n) for n in re.findall(r"\d+", locator)}
    if wanted:
        by_number = [
            b for b in report.blocks
            if {int(n) for n in re.findall(r"\d+", b.locator)} & wanted
        ]
        if by_number:
            return by_number

    if strict:
        return []

    # 都对不上就退回整个文件：粗糙，但总比什么都取不到强
    return list(report.blocks)


def _declared_sources(sources: list | None) -> list[dict]:
    """把 `sources` 里**真正声明了文件**的项挑出来（脏数据一律忽略）。"""
    if not isinstance(sources, list):
        return []
    declared: list[dict] = []
    for item in sources:
        if not isinstance(item, dict):
            continue
        name = str(item.get("file") or "").strip()
        if not name:
            continue
        declared.append({"name": name, "locator": str(item.get("locator") or "").strip()})
    return declared


def material_blocks(
    reports: list[FileReport],
    sources: list | None,
    *,
    query: str = "",
    max_blocks: int = MAX_EXCERPT_BLOCKS,
    strict: bool = False,
) -> list[Block]:
    """挑出**这一轮**要读的资料块：先按 `sources` 定位范围，再按 `query` 精选内容。

    范围这一步是三级退化来的：`file` + `locator` 都命中就只取那一块；
    只有文件名对得上（或 locator 给的是范围）就取该文件里相关的几块；
    全都对不上就退回"整批资料的开头"。**每一步都尽量给出点东西**，
    因为在讲师那一轮，"没有资料可依据"比"依据得不够精确"糟得多。

    **但那个取舍只适用于参考模式。** `strict=True` 时反过来：
    声明了来源却一块都对不上，就**返回空**——上层据此明确告诉用户
    "这一轮没取到可核验的原文"，而不是拿别的材料凑数。
    这条边界是严格模式的意义所在（评审 4.1）。

    精选这一步见 `select_blocks`——它是"按需取用"的落点。
    """
    usable = [r for r in reports if r.ok and r.has_text]
    if not usable:
        return []

    declared = _declared_sources(sources)

    picked: list[Block] = []
    for item in declared:
        matched_any = False
        for report in usable:
            # 文件名用**双向包含**匹配：模型偶尔会多写或少写扩展名，
            # 而旧格式会把页码一起塞在 file 里（"矩阵引论.pdf 第 2 页"）。
            name = item["name"]
            if name not in report.path.name and report.path.name not in name:
                continue
            blocks = _blocks_for(report, item["locator"], strict=strict)
            if blocks:
                matched_any = True
                picked.extend(blocks)
        if strict and not matched_any:
            # 严格模式：一条声明的来源定位不到，就当作"这条来源不可核验"。
            # 继续用别的来源会把"我讲过第 2 页"变成假话。
            continue

    if not picked and (not strict or not declared):
        # 两种情况下可以放宽：
        #   · 参考模式：宁可粗，也要有东西可讲；
        #   · 严格模式但**这一轮没有声明任何具体来源**：整批资料仍然属于
        #     "学生的资料"，用它不算违反"只依据资料"。
        #     真正不能放宽的是"声明了第 2 页却核验不到"——那才是可核验性问题。
        picked = [b for r in usable for b in r.blocks]

    if not picked:
        return []

    return select_blocks(picked, query, max_blocks=max_blocks)


def render_blocks(blocks: list[Block], *, limit: int = 4000) -> str:
    """把挑好的块拼成提示词里那段原文，每块前面带可回指的出处。"""
    lines: list[str] = []
    used = 0
    for block in blocks:
        if used >= limit:
            break
        piece = block.text[: limit - used]
        lines.append(f"[{block.source} {block.locator}]\n{piece}")
        used += len(piece)
    return "\n\n".join(lines)


def material_excerpt(
    reports: list[FileReport],
    sources: list | None,
    *,
    limit: int = 4000,
    query: str = "",
    max_blocks: int = MAX_EXCERPT_BLOCKS,
) -> str:
    """按 `sources` 取回相关资料原文，供讲师这一轮引用（`material_blocks` + 拼文本）。"""
    return render_blocks(
        material_blocks(reports, sources, query=query, max_blocks=max_blocks),
        limit=limit,
    )


# ---------------- 按需取用：从"相关的块"里挑"这一轮要用的块" ----------------

# 切词用：连续的小写字母数字算一个词，每个汉字算一个单位。
# 中文没有空格，所以真正的"词"由下面的二元组承担——见 `_grams`。
_TOKEN = re.compile(r"[a-z0-9]+|[\u4e00-\u9fff]")


def _grams(text: str) -> Counter:
    """把文本压成可比较的"词袋"。

    中文按**字符二元组**（"行列""列与"）当词：零依赖，对中文关键词的区分度够用，
    也不必预装分词库——这个项目对新增依赖一向谨慎。英文与数字按整词。
    单字也留着，但真正起作用的是二元组：单个"行"字在任何一页里都出现过。
    """
    tokens = _TOKEN.findall((text or "").lower())
    grams: Counter = Counter(tokens)
    for first, second in zip(tokens, tokens[1:]):
        grams[first + second] += 1
    return grams


def _relevance(block_grams: Counter, query_grams: Counter, df: Counter, total: int) -> float:
    """块与这一轮语境的贴合度。

    用 df（有多少块含这个 gram）做加权：越是"到处都是"的 gram（"第2页""第一章"）
    越没有区分度，加分越少。没有这一步的话，块越长越容易赢——
    而我们要的是"最对得上"，不是"最长"。
    """
    score = 0.0
    for gram, weight in query_grams.items():
        hit = block_grams.get(gram, 0)
        if not hit:
            continue
        idf = math.log(1 + total / (1 + df.get(gram, 0)))
        # 同一 gram 在语境里出现多次说明它是重点，给一点额外权重，但封顶，
        # 免得一个高频词把其余线索全压过去
        score += min(2, weight) * idf
    # 对长块轻微惩罚：否则"什么都沾一点"的长块总能压过精准的短块
    return score / (1 + 0.02 * math.log(1 + len(block_grams)))


def select_blocks(
    blocks: list[Block], query: str, *, max_blocks: int = 0
) -> list[Block]:
    """在候选块里挑出这一轮真正要读的几块——**"按需取用"就落在这一步**。

    候选范围由 `sources` 定（这个知识点出自哪几页），具体取哪几块由**这一轮的语境**定
    （现在要讲的那一步的关键点 + 学生刚问的那句话）。所以同一个知识点，
    讲第 1 步和第 3 步取到的原文可以不是同一页：

      · 需要学生资料里的原话、例子、定义时才取来，且只取对得上的那几块；
      · 不像"每轮都塞整份资料"那样把上下文撑满，也不像"只许看固定一页"那样够不着。

    `query` 为空、或与资料毫无字面重合时，退回原顺序的前几块：
    **"取不准"的代价远小于"取不到"**（这条与 `material_blocks` 的三级退化同一个原则）。
    """
    if not blocks or max_blocks <= 0 or len(blocks) <= max_blocks:
        return list(blocks)

    query_grams = _grams(query)
    if not query_grams:
        return blocks[:max_blocks]

    block_grams = [_grams(block.text) for block in blocks]
    df: Counter = Counter()
    for grams in block_grams:
        df.update(grams.keys())

    total = len(blocks)
    scored = [
        (_relevance(grams, query_grams, df, total), index)
        for index, grams in enumerate(block_grams)
    ]
    # 按贴合度降序；同分时保持资料原本的先后（否则同一份资料每次取到的块都可能不同，
    # 结果不可复现）
    scored.sort(key=lambda item: (-item[0], item[1]))
    chosen = [index for score, index in scored[:max_blocks] if score > 0]
    if not chosen:
        return blocks[:max_blocks]
    return [blocks[index] for index in sorted(chosen)]


# 讲师提示词里的两句兜底文案。
#
# 抽成常量是因为**文案本身就是行为边界**：散在两处就会变成"一处改了、另一处还是旧的"，
# 而这种不一致的后果很具体——讲师会在没有原文的情况下讲得像"照资料讲的"。
NO_MATERIALS = "（这次没有指定资料，按你自己的知识讲。）"
NO_PASSAGE = (
    "（这一轮没取到资料原文——资料里没有对应段落，或者资料读不到。"
    "按你自己的知识讲，但**不要声称它来自学生的资料**。）"
)
#: 严格模式下"声明的来源一块都没核验上"时的提示词。
#:
#: 与 `NO_PASSAGE` 分开写，是因为两者要讲师做的事**相反**：
#: 前者允许"用你自己的知识讲（但别冒充资料）"，这里则要求
#: "先说明没取到原文，再问学生是否允许补充"——严格模式的学生选的就是这个。
NO_VERIFIED_PASSAGE = (
    "（**严格模式**：这一轮声明的资料出处都没有核验到原文，"
    "所以本轮**没有可依据的资料片段**。请先明确告诉学生「这一段我没取到原文」，"
    "再问是否允许你按自己的知识补充。**在你得到同意之前，不要展开讲解**，"
    "也不要把下面任何内容说成来自他的资料。）"
)


@dataclass
class MaterialView:
    """讲师这一轮的资料视图：提示词那一段 + 这一轮的实际出处。

    `labels` 是给界面用的（"依据：矩阵引论.pdf 第 2 页"）：
    学生据此核对"你讲的是不是我那份资料"——这正是选严格模式的人的诉求。
    两样东西一起产出，是为了避免"取一遍原文、再算一遍出处"这种会走偏的实现
    （两处逻辑一旦分家，界面显示的出处就可能不是讲师真正看的那几页）。

    `strict_miss` 是给界面用的**告警位**：严格模式下声明了来源却一块都没核验上。
    调用方据此推一条 `alert`（不自动消失），让学生当场知道
    "这次讲的没有资料依据"，而不是等他事后去核对才发现。
    """

    note: str
    labels: list[str]
    strict_miss: bool = False


def source_labels(blocks: list[Block], *, limit: int = 2) -> list[str]:
    """这一轮用到的出处，做成给界面看的短标签（去重、最多几个）。"""
    labels: list[str] = []
    for block in blocks:
        label = f"{block.source} {block.locator}".strip()
        if label and label not in labels:
            labels.append(label)
        if len(labels) >= limit:
            break
    return labels


def _note_text(excerpt: str, mode: str) -> str:
    """把原文片段包成"这一轮的行为边界"。

    两种模式给的约束**方向相反**，所以这段话也完全不同。
    它不是背景介绍，而是这一轮的行为边界——
    严格模式下"不许补充"和参考模式下"补了要说明"，是不能互相折中的。

    两种模式共用的一段是**"这些只是挑出来的片段"**：
    不说清，讲师会把手上这几块当成资料的全部，于是把"片段里没有"
    当成"你资料里没有"讲出来——严格模式下这等于骗人。
    """
    picking = (
        "下面这些是**按这一轮要讲的要点和你刚问的问题，从资料里挑出来的片段**，"
        "不是资料的全貌。片段里没提到的，只说「这一段原文我没取到」，"
        "**不要断言「你的资料里没有」**。\n\n"
    )
    if normalize_mode(mode) == MODE_STRICT:
        return (
            "学生要求**严格按资料**讲。\n\n"
            + picking
            + f"本轮相关资料原文：\n\n{excerpt}\n\n"
            "**这一轮必须严格依据上面的资料**：不要补充资料里没有的概念、结论或例子。"
            "如果学生问到这些片段没覆盖的地方，就说「这部分我这边没取到原文」，"
            "再问他是否愿意让你按自己的知识补充——**不要擅自展开**。"
        )
    return (
        "以下是本轮相关资料原文（学生允许你在它基础上补充）。\n\n"
        + picking
        + f"{excerpt}\n\n"
        "**资料讲过的地方，以资料为准**（包括它的术语、顺序与说法）。"
        "资料没覆盖的地方你可以按自己的知识补充，但要让学生分得清"
        "「这是你资料里的」和「这是我补充的」。"
    )


def build_material_view(
    reports: list[FileReport],
    sources: list | None,
    mode: str,
    *,
    query: str = "",
    limit: int = 4000,
    max_blocks: int = MAX_EXCERPT_BLOCKS,
) -> MaterialView:
    """取这一轮的资料原文与出处，并给出对应的行为边界（授课装配用这个）。

    严格模式下的关键差异（评审 4.1）：**声明了来源、但一块都没核验上时，
    不给任何原文**。旧行为是"退回整批资料"——学生指定第 2 页，讲师拿到整份讲义，
    却仍被告知"严格依据资料"，于是它引用到学生没指定的页，
    而界面上显示的出处还是"第 2 页"。那种不一致比"讲不出来"糟得多。
    """
    if not reports:
        return MaterialView(note=NO_MATERIALS, labels=[])

    strict = normalize_mode(mode) == MODE_STRICT
    declared = _declared_sources(sources)

    blocks = material_blocks(
        reports, sources, query=query, max_blocks=max_blocks, strict=strict
    )

    if not blocks:
        if strict and declared:
            # 有声明的来源却核验不到：明确告诉用户"这一轮没有依据"
            return MaterialView(note=NO_VERIFIED_PASSAGE, labels=[], strict_miss=True)
        return MaterialView(note=NO_PASSAGE, labels=[])

    return MaterialView(
        note=_note_text(render_blocks(blocks, limit=limit), mode),
        labels=source_labels(blocks),
    )


def build_materials_note(
    reports: list[FileReport],
    sources: list | None,
    mode: str,
    *,
    query: str = "",
    limit: int = 4000,
) -> str:
    """拼出讲师提示词里「本次可依据的资料」那一段。

    `query` 是"这一轮在讲什么"（当前步骤的要点 + 学生刚问的话）：
    给了它就只挑对得上的那几块原文，见 `select_blocks`。
    """
    return build_material_view(reports, sources, mode, query=query, limit=limit).note
